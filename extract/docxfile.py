"""docx：段落 + 表格一起抽。表格是重点（很多公文/表格模板正文全在表里）。"""
from pathlib import Path

from .common import clean_text

try:
    import docx
except ImportError:
    docx = None


def _table_rows(table) -> list[str]:
    """把一张表拍成文本：一行 = 表格的一行，单元格用 ' | ' 连接。

    ★ 必须按底层元素去重：python-docx 的 row.cells 是"按网格列"返回的，
      **合并单元格会把同一个格子返回好几次**，不去重就会出现
      "姓名 | 姓名 | 王嘉煜" 这种重复，字数直接虚高好几倍。
    """
    out = []
    for tr in table.rows:
        cells, seen = [], set()
        for tc in tr.cells:
            key = id(tc._tc)                 # 同一个合并格 → 同一个底层元素
            if key in seen:
                continue
            seen.add(key)
            text = clean_text(tc.text)
            if text:
                cells.append(text)
        line = " | ".join(cells)
        if line:
            out.append(line)
    return out


def extract_docx(path: Path, doc_id: str = "", rel: str = "") -> list[dict]:
    path = Path(path)
    name, rows = path.name, []
    if docx is None:
        return [{"doc_id": doc_id, "source": name, "path": rel, "type": "docx",
                 "page": 0, "text": "", "missing": True, "note": "未安装python-docx"
                 }]
    try:
        d = docx.Document(str(path))
    except Exception as exc:
        return [{"doc_id": doc_id, "source": name, "path": rel, "type": "docx",
                 "page": 0, "text": "", "missing": True, "note": f"无法打开{type(exc).__name__}: {exc}"
                 }]

    # 正文段落
    for i, p in enumerate(d.paragraphs, start=1):
        text = clean_text(p.text)
        if text:
            rows.append({"doc_id": doc_id, "source": name, "path": rel,
                         "type": "docx", "page": f"p{i}", "text": text, "missing": False, "note": ""})
    # 表格
    for ti, table in enumerate(d.tables, start=1):
        for li, line in enumerate(_table_rows(table), start=1):
            rows.append({"doc_id": doc_id, "source": name, "path": rel,
                         "type": "docx", "page": f"t{ti}r{li}", "text": line, "missing": False, "note": ""})


    # 页眉页脚
    for si, section in enumerate(d.sections, start=1):
        for hf, tag in ((section.header, "header"), (section.footer, "footer")):
            text=clean_text(" ".join(p.text for p in hf.paragraphs))
            if text:
                rows.append({"doc_id": doc_id, "source": name, "path": rel,
                             "type":"docx","page":f"s{si}{tag}","text":text,"missing": False, "note": ""})
    if not rows:
        rows.append({"doc_id": doc_id, "source": name, "path": rel,
                     "type":"docx","page":0,"text": "", "missing": True, "note": "文档里没有可提取文字（可能是图片版）"})
    return rows
